from io import BytesIO
import struct
import unittest
from unittest.mock import patch
from zipfile import ZipFile
import zlib

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import (
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
)

from knowgrain.parsers import DocumentParseError, parse_document


def make_pdf_with_blank_first_page() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    page = writer.add_blank_page(width=200, height=200)

    font = DictionaryObject()
    font[NameObject("/Type")] = NameObject("/Font")
    font[NameObject("/Subtype")] = NameObject("/Type1")
    font[NameObject("/BaseFont")] = NameObject("/Helvetica")
    font_reference = writer._add_object(font)
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_reference}),
        }
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 20 100 Td (Page two text) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)

    result = BytesIO()
    writer.write(result)
    return result.getvalue()


def make_docx_with_paragraph_and_table() -> bytes:
    document = Document()
    document.add_heading("Project", level=1)
    document.add_paragraph("A useful paragraph.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Key"
    table.cell(0, 1).text = "Value"
    result = BytesIO()
    document.save(result)
    return result.getvalue()


def make_docx_with_forged_document_size() -> tuple[bytes, int]:
    valid_package = make_docx_with_paragraph_and_table()
    document_part = "word/document.xml"
    with ZipFile(BytesIO(valid_package), mode="r") as source:
        entries = [
            (info.filename, source.read(info), info.compress_type)
            for info in source.infolist()
        ]

    document_xml = next(data for name, data, _ in entries if name == document_part)
    test_member_limit = max(len(data) for _, data, _ in entries) + 128
    forged_output = BytesIO()
    with ZipFile(forged_output, mode="w") as target:
        for name, data, compression in entries:
            if name == document_part:
                data += b" " * (test_member_limit + 64 * 1024)
            target.writestr(name, data, compress_type=compression)

    forged_package = bytearray(forged_output.getvalue())
    forged_crc = zlib.crc32(
        document_xml + b" " * (test_member_limit + 1 - len(document_xml))
    )
    target_name = document_part.encode("ascii")
    with ZipFile(BytesIO(forged_package), mode="r") as archive:
        info = archive.getinfo(document_part)
        # The compressed stream contains valid XML plus a large whitespace tail;
        # the headers claim the XML size and a forged CRC for output beyond it.
        struct.pack_into("<I", forged_package, info.header_offset + 14, forged_crc)
        struct.pack_into("<I", forged_package, info.header_offset + 22, len(document_xml))

        position = archive.start_dir
        found = False
        while position < len(forged_package):
            if forged_package[position : position + 4] != b"PK\x01\x02":
                break
            filename_size, extra_size, comment_size = struct.unpack_from(
                "<HHH", forged_package, position + 28
            )
            name_start = position + 46
            name_end = name_start + filename_size
            if forged_package[name_start:name_end] == target_name:
                struct.pack_into("<I", forged_package, position + 16, forged_crc)
                struct.pack_into("<I", forged_package, position + 24, len(document_xml))
                found = True
                break
            position = name_end + extra_size + comment_size
        if not found:
            raise AssertionError("DOCX central-directory entry was not found")

    return bytes(forged_package), test_member_limit


class ParseDocumentTests(unittest.TestCase):
    def test_utf8_markdown_preserves_heading_locations(self) -> None:
        parsed = parse_document(
            "notes.MD", "# 项目\n\n第一段。\n\n## 决定\n第二段。".encode("utf-8")
        )

        self.assertEqual(parsed.parser_version, "1")
        self.assertIn("第一段。", parsed.text)
        self.assertEqual(parsed.segments[1].heading, "项目")
        self.assertEqual(parsed.segments[2].heading, "决定")

    def test_invalid_utf8_is_rejected(self) -> None:
        with self.assertRaises(DocumentParseError):
            parse_document("notes.txt", b"\xff\xfe\x81")

    def test_unsupported_and_empty_documents_are_rejected(self) -> None:
        with self.assertRaises(DocumentParseError):
            parse_document("archive.zip", b"not empty")
        with self.assertRaises(DocumentParseError):
            parse_document("empty.txt", b" \n\t ")

    def test_pdf_skips_blank_pages_and_preserves_original_page_numbers(self) -> None:
        parsed = parse_document("source.pdf", make_pdf_with_blank_first_page())

        self.assertEqual(parsed.text, "Page two text")
        self.assertEqual([segment.page for segment in parsed.segments], [2])
        self.assertTrue(all(segment.heading is None for segment in parsed.segments))

    def test_blank_or_scanned_pdf_is_rejected(self) -> None:
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        result = BytesIO()
        writer.write(result)

        with self.assertRaises(DocumentParseError):
            parse_document("blank.pdf", result.getvalue())

    def test_encrypted_pdf_is_rejected(self) -> None:
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        writer.encrypt("secret")
        result = BytesIO()
        writer.write(result)

        with self.assertRaises(DocumentParseError):
            parse_document("locked.pdf", result.getvalue())

    def test_docx_extracts_paragraphs_and_table_rows(self) -> None:
        parsed = parse_document("report.docx", make_docx_with_paragraph_and_table())

        self.assertEqual(
            parsed.text,
            "Project\n\nA useful paragraph.\n\nKey | Value",
        )
        self.assertEqual(parsed.segments[1].heading, "Project")
        self.assertEqual(parsed.segments[2].text, "Key | Value")

    def test_malformed_docx_is_rejected(self) -> None:
        with self.assertRaises(DocumentParseError):
            parse_document("broken.docx", b"not a zip package")

    def test_docx_decompression_bound_is_enforced(self) -> None:
        with patch("knowgrain.parsers._DOCX_MAX_UNCOMPRESSED_BYTES", 1):
            with self.assertRaises(DocumentParseError):
                parse_document("large.docx", make_docx_with_paragraph_and_table())

    def test_docx_forged_small_member_size_cannot_bypass_actual_read_bound(self) -> None:
        package, member_limit = make_docx_with_forged_document_size()

        with patch("knowgrain.parsers._DOCX_MAX_MEMBER_BYTES", member_limit):
            with self.assertRaisesRegex(DocumentParseError, "member size limit"):
                parse_document("forged.docx", package)

    def test_docx_member_count_bound_rejects_many_empty_entries(self) -> None:
        package = BytesIO()
        with ZipFile(package, mode="w") as archive:
            for index in range(40):
                archive.writestr(f"empty-{index}.xml", b"")

        with patch("knowgrain.parsers._DOCX_MAX_ZIP_MEMBERS", 32):
            with self.assertRaises(DocumentParseError):
                parse_document("many-members.docx", package.getvalue())


if __name__ == "__main__":
    unittest.main()
