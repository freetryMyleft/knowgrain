"""Local extraction of supported source document formats."""

from __future__ import annotations

import copy
from dataclasses import dataclass
from io import BytesIO
from pathlib import PurePath
import re
import struct
from zipfile import BadZipFile, ZIP_DEFLATED, ZIP_STORED, ZipFile, ZipInfo


class DocumentParseError(ValueError):
    """A document is unsupported, invalid, encrypted, or has no extractable text."""


@dataclass(slots=True)
class ParsedSegment:
    text: str
    page: int | None = None
    heading: str | None = None


@dataclass(slots=True)
class ParsedDocument:
    text: str
    segments: list[ParsedSegment]
    parser_version: str = "1"


_DOCX_MAX_UNCOMPRESSED_BYTES = 64 * 1024 * 1024
_DOCX_MAX_MEMBER_BYTES = 32 * 1024 * 1024
_DOCX_MAX_COMPRESSION_RATIO = 200
_DOCX_MAX_ZIP_MEMBERS = 4096
_DOCX_READ_CHUNK_BYTES = 64 * 1024


def parse_document(filename: str, content: bytes) -> ParsedDocument:
    """Extract text and available page or heading positions from supported files."""
    suffix = PurePath(filename).suffix.casefold()
    if suffix in {".md", ".markdown", ".txt"}:
        return _parse_text(content, markdown=suffix in {".md", ".markdown"})
    if suffix == ".pdf":
        return _parse_pdf(content)
    if suffix == ".docx":
        return _parse_docx(content)
    raise DocumentParseError("unsupported document type")


def _parse_text(content: bytes, *, markdown: bool) -> ParsedDocument:
    try:
        text = content.decode("utf-8-sig").strip()
    except UnicodeDecodeError as exc:
        raise DocumentParseError("text document is not valid UTF-8") from exc
    if not text:
        raise DocumentParseError("document contains no extractable text")

    segments: list[ParsedSegment] = []
    current_heading: str | None = None
    for block in re.split(r"\n\s*\n", text):
        block = block.strip()
        if not block:
            continue
        if markdown:
            heading_match = re.search(r"(?m)^\s{0,3}#{1,6}\s+(.+?)\s*#*\s*$", block)
            if heading_match:
                current_heading = heading_match.group(1).strip()
        segments.append(ParsedSegment(text=block, heading=current_heading if markdown else None))

    return ParsedDocument(text=text, segments=segments)


def _parse_pdf(content: bytes) -> ParsedDocument:
    try:
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(content), strict=False)
        if reader.is_encrypted:
            raise DocumentParseError("encrypted PDF documents are not supported")
        segments: list[ParsedSegment] = []
        for page_number, page in enumerate(reader.pages, start=1):
            page_text = (page.extract_text() or "").strip()
            # Preserve source page numbering while omitting pages with no text.
            if page_text:
                segments.append(ParsedSegment(text=page_text, page=page_number))
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError("could not read PDF document") from exc

    if not segments:
        raise DocumentParseError("PDF contains no extractable text (it may be scanned)")
    return ParsedDocument(
        text="\n\n".join(segment.text for segment in segments),
        segments=segments,
    )


def _parse_docx(content: bytes) -> ParsedDocument:
    try:
        # ZipFile eagerly builds a ZipInfo object for every central-directory
        # entry. Check the declared entry count from the small end record first
        # so a package with many empty files is rejected before that allocation.
        if _declared_zip_member_count(content) > _DOCX_MAX_ZIP_MEMBERS:
            raise DocumentParseError("DOCX contains too many package entries")

        with ZipFile(BytesIO(content)) as archive:
            if len(archive.infolist()) > _DOCX_MAX_ZIP_MEMBERS:
                raise DocumentParseError("DOCX contains too many package entries")
            total_uncompressed = 0
            for member in archive.infolist():
                total_uncompressed += member.file_size
                if member.file_size > _DOCX_MAX_MEMBER_BYTES:
                    raise DocumentParseError("DOCX expands beyond the supported size limit")
                if (
                    member.file_size > 1024 * 1024
                    and member.compress_size > 0
                    and member.file_size / member.compress_size > _DOCX_MAX_COMPRESSION_RATIO
                ):
                    raise DocumentParseError("DOCX compression ratio exceeds the supported limit")
                if total_uncompressed > _DOCX_MAX_UNCOMPRESSED_BYTES:
                    raise DocumentParseError("DOCX expands beyond the supported size limit")
            bounded_package = _normalize_docx_package(archive)

        from docx import Document
        from docx.oxml.text.paragraph import CT_P
        from docx.oxml.table import CT_Tbl
        from docx.text.paragraph import Paragraph
        from docx.table import Table

        # The original ZIP metadata is untrusted. Pass python-docx only a package
        # rebuilt from payloads whose actual decompressed sizes were bounded.
        document = Document(bounded_package)
        segments: list[ParsedSegment] = []
        current_heading: str | None = None
        for child in document.element.body.iterchildren():
            if isinstance(child, CT_P):
                paragraph = Paragraph(child, document)
                text = paragraph.text.strip()
                if not text:
                    continue
                if paragraph.style is not None and (
                    paragraph.style.name.startswith("Heading")
                    or paragraph.style.name == "Title"
                ):
                    current_heading = text
                segments.append(ParsedSegment(text=text, heading=current_heading))
            elif isinstance(child, CT_Tbl):
                table = Table(child, document)
                for row in table.rows:
                    cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    row_text = " | ".join(cells).strip()
                    if row_text and any(cells):
                        segments.append(ParsedSegment(text=row_text, heading=current_heading))
    except DocumentParseError:
        raise
    except (BadZipFile, OSError, ValueError, KeyError, TypeError) as exc:
        raise DocumentParseError("could not read DOCX document") from exc
    except Exception as exc:
        # python-docx can raise several XML/zip implementation exceptions for a
        # malformed package; expose one stable, user-safe error type.
        raise DocumentParseError("could not read DOCX document") from exc

    if not segments:
        raise DocumentParseError("document contains no extractable text")
    return ParsedDocument(
        text="\n\n".join(segment.text for segment in segments),
        segments=segments,
    )


def _normalize_docx_package(archive: ZipFile) -> BytesIO:
    """Rebuild a DOCX while bounding actual decompressed bytes per member and total.

    ZipExtFile.read() without a size can decompress far beyond ZipInfo.file_size
    before truncating its returned value. Give each reader a bounded sentinel size
    and always read fixed-size chunks, then verify its actual length against the
    original directory entry before making the normalized package available.
    """
    normalized = BytesIO()
    actual_total = 0
    with ZipFile(normalized, mode="w", compression=ZIP_STORED) as safe_archive:
        for member in archive.infolist():
            if member.compress_type not in {ZIP_STORED, ZIP_DEFLATED}:
                raise DocumentParseError("DOCX uses an unsupported ZIP compression method")

            # Ignore an attacker-supplied small size while streaming, but stop as
            # soon as one byte beyond the configured member bound is observed.
            bounded_info = copy.copy(member)
            bounded_info.file_size = _DOCX_MAX_MEMBER_BYTES + 1
            actual_member = 0
            safe_info = ZipInfo(member.filename, date_time=member.date_time)
            safe_info.compress_type = ZIP_STORED
            safe_info.create_system = member.create_system
            safe_info.external_attr = member.external_attr
            safe_info.internal_attr = member.internal_attr

            try:
                with archive.open(bounded_info, mode="r") as source:
                    with safe_archive.open(safe_info, mode="w") as destination:
                        while chunk := source.read(_DOCX_READ_CHUNK_BYTES):
                            actual_member += len(chunk)
                            actual_total += len(chunk)
                            if actual_member > _DOCX_MAX_MEMBER_BYTES:
                                raise DocumentParseError(
                                    "DOCX expands beyond the supported member size limit"
                                )
                            if actual_total > _DOCX_MAX_UNCOMPRESSED_BYTES:
                                raise DocumentParseError(
                                    "DOCX expands beyond the supported size limit"
                                )
                            destination.write(chunk)
            except DocumentParseError:
                raise
            except Exception as exc:
                raise DocumentParseError("could not read DOCX document") from exc

            if actual_member != member.file_size:
                raise DocumentParseError("DOCX package member size does not match its directory")

    normalized.seek(0)
    return normalized


def _declared_zip_member_count(content: bytes) -> int:
    """Bound and validate ZIP central-directory entries before ZipFile loads them."""
    minimum_offset = max(0, len(content) - (22 + 65535))
    marker = b"PK\x05\x06"
    offset = content.rfind(marker, minimum_offset)
    while offset >= minimum_offset:
        if offset + 22 <= len(content):
            (
                _signature,
                disk_number,
                directory_disk,
                disk_entries,
                entries_total,
                directory_size,
                directory_offset,
                comment_length,
            ) = struct.unpack_from("<4s4H2LH", content, offset)
            if offset + 22 + comment_length == len(content):
                if disk_number or directory_disk or disk_entries != entries_total:
                    raise DocumentParseError("multi-disk DOCX packages are not supported")
                # 0xffff is the ZIP64 sentinel and is already above our cap.
                if entries_total > _DOCX_MAX_ZIP_MEMBERS:
                    return entries_total

                # ZIP permits prepended self-extracting data. Derive its length
                # from the directory offsets, then scan entries without creating
                # ZipInfo objects. The count from EOCD alone is untrusted input.
                archive_prefix = offset - directory_size - directory_offset
                directory_start = archive_prefix + directory_offset
                directory_end = directory_start + directory_size
                if (
                    archive_prefix < 0
                    or directory_start < 0
                    or directory_end != offset
                ):
                    raise DocumentParseError("could not read DOCX package directory")

                position = directory_start
                actual_entries = 0
                while position < directory_end:
                    if content[position : position + 4] == b"PK\x05\x05":
                        # The optional central-directory digital signature is
                        # not a member entry and follows all central headers.
                        if position + 6 > directory_end:
                            raise DocumentParseError("could not read DOCX package directory")
                        signature_size = struct.unpack_from("<H", content, position + 4)[0]
                        if position + 6 + signature_size != directory_end:
                            raise DocumentParseError("could not read DOCX package directory")
                        position = directory_end
                        break
                    if (
                        position + 46 > directory_end
                        or content[position : position + 4] != b"PK\x01\x02"
                    ):
                        raise DocumentParseError("could not read DOCX package directory")
                    filename_size, extra_size, member_comment_size = struct.unpack_from(
                        "<HHH", content, position + 28
                    )
                    position += 46 + filename_size + extra_size + member_comment_size
                    if position > directory_end:
                        raise DocumentParseError("could not read DOCX package directory")
                    actual_entries += 1
                    if actual_entries > _DOCX_MAX_ZIP_MEMBERS:
                        return actual_entries

                if actual_entries != entries_total:
                    raise DocumentParseError("could not read DOCX package directory")
                return actual_entries
        offset = content.rfind(marker, minimum_offset, offset)
    raise DocumentParseError("could not read DOCX package directory")
